import os
import csv
import math
import traceback
from io import BytesIO, StringIO
from datetime import datetime, date
from pathlib import Path
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from sqlalchemy import create_engine, select, or_, func
from sqlalchemy.exc import IntegrityError, DataError, StatementError
from sqlalchemy.orm import sessionmaker, joinedload
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from decimal import Decimal, ROUND_HALF_UP
from models import (
    Base,
    Invoice,
    Trip,
    MonthlyBill,
    MonthlyBillItem,
)
from services import (
    calculate_invoice,
    build_invoice_pdf,
    invoice_to_csv,
    invoice_to_excel,
    invoice_to_einvoice_excel,
)

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR.parent / ".env")
load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL")
if not DATABASE_URL:
    raise RuntimeError("DATABASE_URL is not configured. Set it in .env locally or Vercel Environment Variables.")

if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
    pool_recycle=300,
    connect_args={"connect_timeout": 10},
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

app = Flask(__name__)
origins = os.getenv("FRONTEND_ORIGIN", "*")
CORS(app, resources={r"/api/*": {"origins": origins.split(",") if origins != "*" else "*"}})


def init_tables():
    Base.metadata.create_all(bind=engine)


# Safe for an existing database: CREATE TABLE IF NOT EXISTS semantics are used by SQLAlchemy.
try:
    init_tables()
except Exception:
    # Do not hide a database problem from the API; requests will expose a useful health/error response.
    traceback.print_exc()


def error_response(message, status=400, code=None, details=None):
    payload = {"success": False, "message": message}
    if code:
        payload["code"] = code
    if details:
        payload["details"] = details
    return jsonify(payload), status


def parse_date(value, field="date"):
    if not value:
        raise ValueError(f"{field} is required")
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        raise ValueError(f"{field} must be a valid date in YYYY-MM-DD format")


def serialize_trip(t):
    return {
        "id": t.id,
        "ds_no": t.ds_no or "",
        "trip_date": t.trip_date.isoformat() if t.trip_date else "",
        "end_date": t.end_date.isoformat() if t.end_date else "",
        "vehicle_type": t.vehicle_type or "",
        "vehicle_number": t.vehicle_number or "",
        "start_time": t.start_time or "",
        "end_time": t.end_time or "",
        "start_km": t.start_km or 0,
        "end_km": t.end_km or 0,
        "total_hours": t.total_hours or 0,
        "total_km": t.total_km or 0,
        "slab_hours": t.slab_hours or 0,
        "slab_km": t.slab_km or 0,
        "slab_rate": t.slab_rate or 0,
        "extra_hour_rate": t.extra_hour_rate or 0,
        "extra_km_rate": t.extra_km_rate or 0,
        "extra_hours": t.extra_hours or 0,
        "extra_km": t.extra_km or 0,
        "extra_hour_amount": t.extra_hour_amount or 0,
        "extra_km_amount": t.extra_km_amount or 0,
        "base_amount": t.base_amount or 0,
        "driver_bata": t.driver_bata or 0,
        "parking": t.parking or 0,
        "toll": t.toll or 0,
        "other_charges": t.other_charges or 0,
        "trip_total": t.trip_total or 0,
        "notes": t.notes or "",
    }


def serialize_invoice(inv):
    return {
        "id": inv.id,
        "invoice_number": inv.invoice_number,
        "invoice_series": inv.invoice_series,
        "invoice_serial_number": inv.invoice_serial_number,
        "invoice_date": inv.invoice_date.isoformat() if inv.invoice_date else "",
        "customer_name": inv.customer_name,
        "customer_address": inv.customer_address or "",
        "customer_gstin": inv.customer_gstin or "",
        "booked_by": inv.booked_by or "",
        "used_by": inv.used_by or "",
        "reference_number": inv.reference_number or "",
        "cgst_rate": inv.cgst_rate or 0,
        "sgst_rate": inv.sgst_rate or 0,
        "igst_rate": inv.igst_rate or 0,
        "subtotal": inv.subtotal or 0,
        "cgst": inv.cgst or 0,
        "sgst": inv.sgst or 0,
        "igst": inv.igst or 0,
        "round_off": inv.round_off or 0,
        "grand_total": inv.grand_total or 0,
        "created_at": inv.created_at.isoformat() if inv.created_at else None,
        "updated_at": inv.updated_at.isoformat() if inv.updated_at else None,
        "trips": [serialize_trip(t) for t in inv.trips],
    }

def serialize_monthly_bill(bill):
    return {
        "id": bill.id,
        "invoice_number": bill.invoice_number,
        "invoice_date": (
            bill.invoice_date.isoformat()
            if bill.invoice_date
            else ""
        ),
        "customer_name": bill.customer_name or "",
        "customer_profile": getattr(bill, "customer_profile", None) or "",
        "customer_address": bill.customer_address or "",
        "customer_gstin": bill.customer_gstin or "",
        "used_by": getattr(bill, "used_by", None) or "",
        "start_date": (
            bill.start_date.isoformat()
            if getattr(bill, "start_date", None)
            else ""
        ),
        "end_date": (
            bill.end_date.isoformat()
            if getattr(bill, "end_date", None)
            else ""
        ),
        "booked_by": bill.booked_by or "",
        "position": getattr(bill, "position", None) or "",
        "vehicle_number": bill.vehicle_number or "",
        
        "total_km_driven": bill.total_km_driven or 0,
        "km_per_liter": bill.km_per_liter or 0,
        "fuel_quantity": getattr(bill, "fuel_quantity", 0) or 0,
        "fuel_rate": getattr(bill, "fuel_rate", 0) or 0,
        "outstation_duties": getattr(bill, "outstation_duties", 0) or 0,
        "extra_hours": getattr(bill, "extra_hours", 0) or 0,
        "holidays_worked": getattr(bill, "holidays_worked", 0) or 0,
        "reference_number": bill.reference_number or "",

        "taxable_subtotal": bill.taxable_subtotal or 0,
        "cgst_rate": bill.cgst_rate or 0,
        "cgst": bill.cgst or 0,
        "sgst_rate": bill.sgst_rate or 0,
        "sgst": bill.sgst or 0,
        "non_taxable_total": bill.non_taxable_total or 0,
        "engine_oil_amount": getattr(bill, "engine_oil_amount", 0) or 0,
        "parking_toll_amount": getattr(bill, "parking_toll_amount", 0) or 0,
        "telephone_charges_amount": getattr(bill, "telephone_charges_amount", 0) or 0,
        "round_off": bill.round_off or 0,
        "grand_total": bill.grand_total or 0,

        "created_at": (
            bill.created_at.isoformat()
            if bill.created_at
            else None
        ),
        "updated_at": (
            bill.updated_at.isoformat()
            if bill.updated_at
            else None
        ),

        "items": [
            {
                "id": item.id,
                "item_type": item.item_type,
                "description": item.description or "",
                "quantity": item.quantity or 0,
                "rate": item.rate or 0,
                "amount": item.amount or 0,
            }
            for item in bill.items
        ],
    }
def validate_payload(data):
    if not isinstance(data, dict):
        raise ValueError("Request body must be a JSON object")
    customer_name = str(data.get("customer_name") or "").strip()
    if not customer_name:
        raise ValueError("Customer Name is required")
    series = str(data.get("invoice_series") or "").strip()
    serial = str(data.get("invoice_serial_number") or "").strip()
    if not series:
        raise ValueError("Invoice Series is required")
    if not serial:
        raise ValueError("Invoice Serial Number is required")
    if not data.get("invoice_date"):
        raise ValueError("Invoice Date is required")
    invoice_date = parse_date(data["invoice_date"], "Invoice Date")
    invoice_number = f"{series}{serial}"
    trips = data.get("trips")
    if not isinstance(trips, list) or not trips:
        raise ValueError("At least one trip is required")
    return invoice_number, invoice_date


def apply_invoice(inv, data, calculated):
    invoice_number, invoice_date = validate_payload(data)
    inv.invoice_number = invoice_number
    inv.invoice_series = str(data.get("invoice_series") or "").strip()
    inv.invoice_serial_number = str(data.get("invoice_serial_number") or "").strip()
    inv.invoice_date = invoice_date
    inv.customer_name = str(data.get("customer_name") or "").strip()
    inv.customer_address = str(data.get("customer_address") or "").strip()
    inv.customer_gstin = str(data.get("customer_gstin") or "").strip()
    inv.booked_by = str(data.get("booked_by") or "").strip()
    inv.used_by = str(data.get("used_by") or "").strip()
    inv.reference_number = str(data.get("reference_number") or "").strip()
    for key in ("cgst_rate", "sgst_rate", "igst_rate", "subtotal", "cgst", "sgst", "igst", "round_off", "grand_total"):
        setattr(inv, key, calculated[key])
    return inv


@app.get("/api/health")
def health():
    try:
        with engine.connect() as conn:
            conn.exec_driver_sql("SELECT 1")
        return jsonify({"success": True, "status": "ok", "database": "connected"})
    except Exception as exc:
        return error_response("Database connection failed", 503, "DATABASE_UNAVAILABLE", str(exc))

@app.get("/api/invoices")
def list_invoices():
    session = SessionLocal()

    try:
        search = (
            request.args.get("search") or ""
        ).strip()

        month_value = (
            request.args.get("month") or ""
        ).strip()

        year_value = (
            request.args.get("year") or ""
        ).strip()

        query = (
            select(Invoice)
            .options(joinedload(Invoice.trips))
            .order_by(
                Invoice.invoice_date.desc(),
                Invoice.id.desc()
            )
        )

        if search:
            term = f"%{search}%"

            query = query.where(
                or_(
                    Invoice.invoice_number.ilike(term),
                    Invoice.customer_name.ilike(term),
                    Invoice.reference_number.ilike(term)
                )
            )

        # Monthly filtering.
        # If month and year are not supplied,
        # all invoices are returned exactly as before.
        if month_value and year_value:
            try:
                selected_month = int(month_value)
                selected_year = int(year_value)

                if selected_month < 1 or selected_month > 12:
                    return error_response(
                        "Month must be between 1 and 12",
                        400,
                        "INVALID_MONTH"
                    )

                if selected_year < 2000 or selected_year > 2100:
                    return error_response(
                        "Invalid year",
                        400,
                        "INVALID_YEAR"
                    )

                start_date = date(
                    selected_year,
                    selected_month,
                    1
                )

                if selected_month == 12:
                    next_month_date = date(
                        selected_year + 1,
                        1,
                        1
                    )
                else:
                    next_month_date = date(
                        selected_year,
                        selected_month + 1,
                        1
                    )

                query = query.where(
                    Invoice.invoice_date >= start_date,
                    Invoice.invoice_date < next_month_date
                )

            except ValueError:
                return error_response(
                    "Month and year must be valid numbers",
                    400,
                    "INVALID_MONTH_YEAR"
                )

        invoices = (
            session.execute(query)
            .unique()
            .scalars()
            .all()
        )

        rows = []

        for inv in invoices:
            rows.append({
                "id": inv.id,
                "invoice_number": inv.invoice_number,
                "invoice_date": (
                    inv.invoice_date.isoformat()
                    if inv.invoice_date
                    else ""
                ),
                "customer_name": inv.customer_name,
                "reference_number": (
                    inv.reference_number or ""
                ),
                "trip_count": len(inv.trips),
                "subtotal": inv.subtotal or 0,
                "grand_total": inv.grand_total or 0,
            })

        return jsonify({
            "success": True,
            "invoices": rows
        })

    finally:
        session.close()

@app.post("/api/invoices")
def create_invoice():
    session = SessionLocal()

    try:
        data = request.get_json(silent=True) or {}

        invoice_number, _ = validate_payload(data)

        existing = session.execute(
            select(Invoice.id).where(
                Invoice.invoice_number == invoice_number
            )
        ).scalar_one_or_none()

        if existing is not None:
            return error_response(
                "Invoice number already exists",
                409,
                "DUPLICATE_INVOICE_NUMBER"
            )

        calculated_trips, totals = calculate_invoice(
            data,
            data["trips"]
        )

        inv = Invoice()

        apply_invoice(
            inv,
            data,
            totals
        )

        session.add(inv)
        session.flush()

        for tdata in calculated_trips:
            trip_date = (
                parse_date(
                    tdata["trip_date"],
                    "Trip Date"
                )
                if tdata.get("trip_date")
                else None
            )

            end_date = (
                parse_date(
                    tdata["end_date"],
                    "End Date"
                )
                if tdata.get("end_date")
                else trip_date
            )

            tdata["trip_date"] = trip_date
            tdata["end_date"] = end_date

            inv.trips.append(
                Trip(**tdata)
            )

        session.commit()

        refreshed = (
            session.execute(
                select(Invoice)
                .options(
                    joinedload(Invoice.trips)
                )
                .where(
                    Invoice.id == inv.id
                )
            )
            .unique()
            .scalar_one()
        )

        return jsonify({
            "success": True,
            "message": "Invoice created successfully",
            "invoice": serialize_invoice(refreshed)
        }), 201

    except ValueError as exc:
        session.rollback()

        return error_response(
            str(exc),
            400,
            "VALIDATION_ERROR"
        )

    except IntegrityError as exc:
        session.rollback()

        return error_response(
            "Invoice number already exists",
            409,
            "DUPLICATE_INVOICE_NUMBER",
            str(getattr(exc, "orig", exc))
        )

    except (DataError, StatementError) as exc:
        session.rollback()

        return error_response(
            "Invalid data type or value sent to PostgreSQL",
            400,
            "DATABASE_DATA_ERROR",
            str(getattr(exc, "orig", exc))
        )

    except Exception as exc:
        session.rollback()
        traceback.print_exc()

        return error_response(
            "Unable to save invoice",
            500,
            "SERVER_ERROR",
            str(exc)
        )

    finally:
        session.close()
@app.get("/api/invoices/<int:invoice_id>")
def get_invoice(invoice_id):
    session = SessionLocal()

    try:
        inv = (
            session.execute(
                select(Invoice)
                .options(joinedload(Invoice.trips))
                .where(Invoice.id == invoice_id)
            )
            .unique()
            .scalar_one_or_none()
        )

        if not inv:
            return error_response(
                "Invoice not found",
                404,
                "NOT_FOUND"
            )

        return jsonify({
            "success": True,
            "invoice": serialize_invoice(inv)
        })

    finally:
        session.close()





@app.put("/api/invoices/<int:invoice_id>")
def update_invoice(invoice_id):
    session = SessionLocal()
    try:
        data = request.get_json(silent=True) or {}
        inv = session.execute(select(Invoice).options(joinedload(Invoice.trips)).where(Invoice.id == invoice_id)).unique().scalar_one_or_none()
        if not inv:
            return error_response("Invoice not found", 404, "NOT_FOUND")
        invoice_number, _ = validate_payload(data)
        other_id = session.execute(select(Invoice.id).where(Invoice.invoice_number == invoice_number, Invoice.id != invoice_id)).scalar_one_or_none()
        if other_id is not None:
            return error_response("Invoice number already exists on another invoice", 409, "DUPLICATE_INVOICE_NUMBER")

        calculated_trips, totals = calculate_invoice(data, data["trips"])
        apply_invoice(inv, data, totals)

        # Full replacement of child rows inside the same transaction is deterministic and prevents stale trips.
        inv.trips.clear()
        session.flush()

        for tdata in calculated_trips:
            trip_date = parse_date(tdata["trip_date"], "Trip Date") if tdata.get("trip_date") else None
            end_date = parse_date(tdata["end_date"], "End Date") if tdata.get("end_date") else trip_date

            tdata["trip_date"] = trip_date
            tdata["end_date"] = end_date

            inv.trips.append(Trip(**tdata))

        session.commit()
        refreshed = session.execute(select(Invoice).options(joinedload(Invoice.trips)).where(Invoice.id == invoice_id)).unique().scalar_one()
        return jsonify({"success": True, "message": "Invoice updated successfully", "invoice": serialize_invoice(refreshed)})
    except ValueError as exc:
        session.rollback()
        return error_response(str(exc), 400, "VALIDATION_ERROR")
    except IntegrityError as exc:
        session.rollback()
        diag = getattr(getattr(exc, "orig", None), "diag", None)
        constraint_name = getattr(diag, "constraint_name", None)
        if constraint_name in ("ix_invoices_invoice_number", "invoices_invoice_number_key"):
            return error_response("Invoice number already exists on another invoice", 409, "DUPLICATE_INVOICE_NUMBER")
        return error_response("Database constraint error while updating the invoice", 400, "DATABASE_CONSTRAINT_ERROR", str(getattr(exc, "orig", exc)))
    except (DataError, StatementError) as exc:
        session.rollback()
        return error_response("Invalid data type or value sent to PostgreSQL", 400, "DATABASE_DATA_ERROR", str(getattr(exc, "orig", exc)))
    except Exception as exc:
        session.rollback()
        traceback.print_exc()
        return error_response("Unexpected server error while updating invoice", 500, "SERVER_ERROR", str(exc))
    finally:
        session.close()


@app.delete("/api/invoices/<int:invoice_id>")
def delete_invoice(invoice_id):
    session = SessionLocal()
    try:
        inv = session.get(Invoice, invoice_id)
        if not inv:
            return error_response("Invoice not found", 404, "NOT_FOUND")
        session.delete(inv)
        session.commit()
        return jsonify({"success": True, "message": "Invoice deleted successfully"})
    except Exception as exc:
        session.rollback()
        return error_response("Unable to delete invoice", 500, "SERVER_ERROR", str(exc))
    finally:
        session.close()


@app.get("/api/invoices/<int:invoice_id>/pdf")
def invoice_pdf(invoice_id):
    session = SessionLocal()
    try:
        inv = session.execute(select(Invoice).options(joinedload(Invoice.trips)).where(Invoice.id == invoice_id)).unique().scalar_one_or_none()
        if not inv:
            return error_response("Invoice not found", 404, "NOT_FOUND")
        pdf = build_invoice_pdf(inv)
        filename = f"{inv.invoice_number.replace('/', '-')}.pdf"
        return send_file(__import__("io").BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=filename)
    except Exception as exc:
        traceback.print_exc()
        return error_response("Unable to generate PDF", 500, "PDF_ERROR", str(exc))
    finally:
        session.close()


@app.get("/api/invoices/export/csv")
def export_csv():
    session = SessionLocal()

    try:
        invoices = (
            session.execute(
                select(Invoice)
                .options(joinedload(Invoice.trips))
                .order_by(
                    Invoice.invoice_date.desc(),
                    Invoice.id.desc()
                )
            )
            .unique()
            .scalars()
            .all()
        )

        csv_bytes = invoice_to_csv(invoices)

        return send_file(
            __import__("io").BytesIO(csv_bytes),
            mimetype="text/csv; charset=utf-8",
            as_attachment=True,
            download_name="pvr-invoices.csv"
        )

    finally:
        session.close()
@app.get("/api/invoices/export/xlsx")
def export_excel():
    session = SessionLocal()

    try:
        invoices = (
            session.execute(
                select(Invoice)
                .options(joinedload(Invoice.trips))
                .order_by(
                    Invoice.invoice_date.desc(),
                    Invoice.id.desc()
                )
            )
            .unique()
            .scalars()
            .all()
        )

        excel_bytes = invoice_to_excel(invoices)

        return send_file(
            __import__("io").BytesIO(excel_bytes),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="PVR_Invoice_Register.xlsx"
        )

    except Exception as exc:
        traceback.print_exc()

        return error_response(
            "Unable to export invoices to Excel",
            500,
            "EXCEL_EXPORT_ERROR",
            str(exc)
        )

    finally:
        session.close()
# =========================================================
# E-INVOICE EXCEL EXPORT
# =========================================================

@app.get("/api/invoices/export/einvoice-xlsx")
def export_einvoice_excel():
    session = SessionLocal()

    try:
        invoices = (
            session.execute(
                select(Invoice)
                .options(
                    joinedload(Invoice.trips)
                )
                .order_by(
                    Invoice.invoice_date.desc(),
                    Invoice.id.desc()
                )
            )
            .unique()
            .scalars()
            .all()
        )

        excel_bytes = invoice_to_einvoice_excel(
            invoices
        )

        return send_file(
            __import__("io").BytesIO(
                excel_bytes
            ),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="PVR_E_Invoice.xlsx"
        )

    except Exception as exc:
        traceback.print_exc()

        return error_response(
            "Unable to export E-Invoice Excel",
            500,
            "EINVOICE_EXCEL_EXPORT_ERROR",
            str(exc)
        )

    finally:
        session.close()
# =========================================================
# MONTHLY BILL CSV EXPORT
# =========================================================


@app.get("/api/monthly-bills/export/csv")
def export_monthly_bills_csv():
    session = SessionLocal()

    try:
        bills = (
            session.execute(
                select(MonthlyBill)
                .options(joinedload(MonthlyBill.items))
                .order_by(
                    MonthlyBill.invoice_date.desc(),
                    MonthlyBill.id.desc(),
                )
            )
            .unique()
            .scalars()
            .all()
        )

        output = StringIO()
        writer = csv.writer(output)

        writer.writerow([
            "Invoice Number",
            "Invoice Date",
            "Customer Name",
            "Customer Profile",
            "Customer Address",
            "Customer GSTIN",
            "Used By",
            "Booked By",
            "Position",
            "Start Date",
            "End Date",
            "Vehicle Number",
            "Reference / PO",
            "Total KM Driven",
            "KM Per Litre",
            "Fuel Quantity",
            "Fuel Rate",
            "Outstation Duties",
            "Extra Hours",
            "Holidays Worked",
            "Charges Before GST",
            "CGST Rate",
            "CGST",
            "SGST Rate",
            "SGST",
            "Round Off",
            "Grand Total",
            "Items",
        ])

        for bill in bills:
            item_descriptions = []

            for item in bill.items:
                item_descriptions.append(
                    f"{item.description or ''} "
                    f"(Qty: {float(item.quantity or 0):g}, "
                    f"Rate: {float(item.rate or 0):,.2f}, "
                    f"Amount: {float(item.amount or 0):,.2f})"
                )

            charges_before_gst = (
                 float(bill.taxable_subtotal or 0)
                   + float(bill.non_taxable_total or 0)
            )

            writer.writerow([
                bill.invoice_number or "",
                bill.invoice_date.isoformat()
                if bill.invoice_date else "",
                bill.customer_name or "",
                getattr(bill, "customer_profile", None) or "",
                bill.customer_address or "",
                bill.customer_gstin or "",
                getattr(bill, "used_by", None) or "",
                bill.booked_by or "",
                getattr(bill, "position", None) or "",
                bill.start_date.isoformat()
                if getattr(bill, "start_date", None) else "",
                bill.end_date.isoformat()
                if getattr(bill, "end_date", None) else "",
                bill.vehicle_number or "",
                bill.reference_number or "",
                f"{float(bill.total_km_driven or 0):.2f}",
                f"{float(bill.km_per_liter or 0):.2f}",
                f"{float(bill.fuel_quantity or 0):.2f}",
                f"{float(bill.fuel_rate or 0):.2f}",
                f"{float(bill.outstation_duties or 0):.2f}",
                f"{float(bill.extra_hours or 0):.2f}",
                f"{float(bill.holidays_worked or 0):.2f}",
                f"{charges_before_gst:.2f}",
                f"{float(bill.cgst_rate or 0):.2f}",
                f"{float(bill.cgst or 0):.2f}",
                f"{float(bill.sgst_rate or 0):.2f}",
                f"{float(bill.sgst or 0):.2f}",
                f"{float(bill.round_off or 0):.2f}",
                f"{float(bill.grand_total or 0):.2f}",
                " | ".join(item_descriptions),
            ])

        csv_bytes = output.getvalue().encode("utf-8-sig")

        return send_file(
            BytesIO(csv_bytes),
            mimetype="text/csv; charset=utf-8",
            as_attachment=True,
            download_name="PVR_Monthly_Bills.csv",
        )

    except Exception as exc:
        traceback.print_exc()
        return error_response(
            "Unable to export Monthly Bills to CSV",
            500,
            "MONTHLY_CSV_EXPORT_ERROR",
            str(exc),
        )

    finally:
        session.close()


# =========================================================
# MONTHLY BILL EXCEL EXPORT
# =========================================================


@app.get("/api/monthly-bills/export/xlsx")
def export_monthly_bills_excel():
    session = SessionLocal()

    try:
        bills = (
            session.execute(
                select(MonthlyBill)
                .options(joinedload(MonthlyBill.items))
                .order_by(
                    MonthlyBill.invoice_date.desc(),
                    MonthlyBill.id.desc(),
                )
            )
            .unique()
            .scalars()
            .all()
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "PVR Monthly Bill"

        headers = [
            "Invoice Number",
            "Invoice Date",
            "Customer Name",
            "Customer Profile",
            "Customer Address",
            "Customer GSTIN",
            "Used By",
            "Booked By",
            "Position",
            "Start Date",
            "End Date",
            "Vehicle Number",
            "Reference / PO",
            "Total KM Driven",
            "KM Per Litre",
            "Fuel Quantity",
            "Fuel Rate",
            "Monthly Hiring",
            "Outstation Duties",
            "Extra Hours",
            "Holidays Worked",
            "Charges Before GST",
            "CGST %",
            "CGST",
            "SGST %",
            "SGST",
            "Engine Oil",
            "Parking & Toll",
            "Telephone Charges",
            "Round Off",
            "Grand Total",
        ]
        sheet.append(headers)

        item_sheet = workbook.create_sheet("Monthly Bill Items")
        item_sheet.append([
            "Invoice Number",
            "Invoice Date",
            "Customer Name",
            "Description",
            "Quantity",
            "Rate",
            "Amount",
        ])

        for bill in bills:
            charges_before_gst = (
                 float(bill.taxable_subtotal or 0)
                    + float(bill.non_taxable_total or 0)
            )

            monthly_hiring = sum(
                float(item.amount or 0)
                for item in bill.items
                if "monthly hiring" in (item.description or "").lower()
            )

            sheet.append([
                bill.invoice_number or "",
                bill.invoice_date or "",
                bill.customer_name or "",
                getattr(bill, "customer_profile", None) or "",
                bill.customer_address or "",
                bill.customer_gstin or "",
                getattr(bill, "used_by", None) or "",
                bill.booked_by or "",
                getattr(bill, "position", None) or "",
                getattr(bill, "start_date", None) or "",
                getattr(bill, "end_date", None) or "",
                bill.vehicle_number or "",
                bill.reference_number or "",
                float(bill.total_km_driven or 0),
                float(bill.km_per_liter or 0),
                float(bill.fuel_quantity or 0),
                float(bill.fuel_rate or 0),
                monthly_hiring,
                float(bill.outstation_duties or 0),
                float(bill.extra_hours or 0),
                float(bill.holidays_worked or 0),
                charges_before_gst,
                float(bill.cgst_rate or 0),
                float(bill.cgst or 0),
                float(bill.sgst_rate or 0),
                float(bill.sgst or 0),
                float(getattr(bill, "engine_oil_amount", 0) or 0),
                float(getattr(bill, "parking_toll_amount", 0) or 0),
                float(getattr(bill, "telephone_charges_amount", 0) or 0),
                float(bill.round_off or 0),
                float(bill.grand_total or 0),
            ])

            for item in bill.items:
                item_sheet.append([
                    bill.invoice_number or "",
                    bill.invoice_date or "",
                    bill.customer_name or "",
                    item.description or "",
                    float(item.quantity or 0),
                    float(item.rate or 0),
                    float(item.amount or 0),
                ])

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="102A43",
        )
        header_font = Font(bold=True, color="FFFFFF")
        thin_side = Side(style="thin", color="D9E2EC")
        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )
        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        for ws in (sheet, item_sheet):
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = border

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 32

        for row in sheet.iter_rows(min_row=2):
            for index in range(14, 32):
                row[index - 1].number_format = "#,##0.00"

        for row in item_sheet.iter_rows(min_row=2):
            for index in (5, 6, 7):
                row[index - 1].number_format = "#,##0.00"

        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12), 35
            )

        for column_cells in item_sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            item_sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12), 40
            )

        excel_buffer = BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)

        return send_file(
            excel_buffer,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="PVR_Monthly_Bills.xlsx",
        )

    except Exception as exc:
        traceback.print_exc()
        return error_response(
            "Unable to export Monthly Bills to Excel",
            500,
            "MONTHLY_EXCEL_EXPORT_ERROR",
            str(exc),
        )

    finally:
        session.close()





def calculate_monthly_bill_data(data):
    """Calculate Monthly Bill charges using the selected customer profile."""

    # ---------------------------------------------------------
    # 1. COMMON VALIDATION HELPERS
    # ---------------------------------------------------------

    def number(key, default=0):
        value = data.get(key)
        value = default if value in (None, "") else value

        try:
            result = float(value)
        except (TypeError, ValueError):
            raise ValueError(
                f"{key.replace('_', ' ').title()} must be a valid number"
            )

        if not math.isfinite(result):
            raise ValueError(
                f"{key.replace('_', ' ').title()} must be a valid number"
            )

        if result < 0:
            raise ValueError(
                f"{key.replace('_', ' ').title()} cannot be negative"
            )

        return result

    def optional_date(key):
        value = data.get(key)

        return (
            parse_date(value, key.replace("_", " ").title())
            if value else None
        )

    # ---------------------------------------------------------
    # 2. CUSTOMER PROFILE AND FIXED RATES
    # ---------------------------------------------------------

    profile = str(data.get("customer_profile") or "").strip()

    profile_key = "".join(
        ch for ch in profile.upper() if ch.isalnum()
    )

    profiles = {
        "SAMRO": {
            "monthly": 63000,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
            "samro": True,
        },
        "DGM": {
            "monthly": 54899,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
        },
        "GM": {
            "monthly": 63000,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
        },
        "CGM": {
            "monthly": 145899,
            "extra": 250,
            "batta": 1000,
            "holiday": 1200,
        },
        "AGM": {
            "monthly": 38899,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
        },
        "GM1": {
            "monthly": 61899,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
        },
        "GTFC": {
            "monthly": 54899,
            "extra": 100,
            "batta": 500,
            "holiday": 500,
        },
        "AOCYBERABAD": {
            "monthly": 36700,
            "extra": 130,
            "batta": 500,
            "holiday": 500,
        },
    }

    # Identify the selected profile.
    if "SAMRO" in profile_key:
        config = profiles["SAMRO"]

    elif "CYBERABAD" in profile_key:
        config = profiles["AOCYBERABAD"]

    elif "CGM" in profile_key:
        config = profiles["CGM"]

    elif "AGM" in profile_key:
        config = profiles["AGM"]

    elif "GM1" in profile_key:
        config = profiles["GM1"]

    elif "DGM" in profile_key:
        config = profiles["DGM"]

    elif "GTFC" in profile_key:
        config = profiles["GTFC"]

    elif profile_key.endswith("GM"):
        config = profiles["GM"]

    else:
        raise ValueError("Please select a valid Customer Profile")

    monthly_hiring = config["monthly"]
    extra_rate = config["extra"]
    batta_rate = config["batta"]
    holiday_rate = config["holiday"]

    is_samro = config.get("samro", False)

    # ---------------------------------------------------------
    # 3. MONTHLY USAGE INPUTS
    # ---------------------------------------------------------

    total_km = number("total_km_driven")
    km_per_liter = number("km_per_liter")
    fuel_rate = number("fuel_rate")

    outstation_duties = number("outstation_duties")
    extra_hours = number("extra_hours")
    holidays_worked = number("holidays_worked")

    engine_oil = number("engine_oil_amount")
    parking_toll = number("parking_toll_amount")
    telephone_charges = number("telephone_charges_amount")

    # ---------------------------------------------------------
    # 4. FUEL CALCULATION
    # ---------------------------------------------------------

    if is_samro:
        # SAMRO:
        # Fuel quantity and fuel rate are entered manually.

        fuel_quantity = number("fuel_quantity")

        fuel_amount = fuel_quantity * fuel_rate

    else:
        # Other profiles:
        # Fuel quantity = Total KM / KM per Litre.

        if total_km > 0 and km_per_liter <= 0:
            raise ValueError(
                "KM per Liter is required when Total KM Driven is entered"
            )

        fuel_quantity = (
            total_km / km_per_liter
            if km_per_liter > 0 else 0
        )

        fuel_amount = fuel_quantity * fuel_rate

    # ---------------------------------------------------------
    # 5. ENGINE OIL CALCULATION
    # ---------------------------------------------------------

    # If the user enters an engine oil amount, use that amount.
    # Otherwise, for non-SAMRO profiles:
    # Engine oil = (Total KM / 250) * 300.

    if not engine_oil and not is_samro:
        engine_oil = (total_km / 250) * 300

    # ---------------------------------------------------------
    # 6. PROFILE-BASED CHARGE CALCULATIONS
    # ---------------------------------------------------------

    # Outstation Batta = Number of Duties * Profile Batta Rate.
    batta_amount = outstation_duties * batta_rate

    # Extra Hours = Number of Extra Hours * Profile Extra Rate.
    extra_amount = extra_hours * extra_rate

    # Holiday Allowance = Holidays Worked * Profile Holiday Rate.
    holiday_amount = holidays_worked * holiday_rate

    # Monthly Hiring is the fixed profile amount.
    monthly_hiring_amount = monthly_hiring

    # ---------------------------------------------------------
    # 7. CREATE BILL LINE ITEMS
    # ---------------------------------------------------------
    # item_type is maintained internally for accounting.
    # The user does not enter or select this classification.

    items = [
        {
            "item_type": "taxable",
            "description": "Monthly Hiring Charges",
            "quantity": 1,
            "rate": monthly_hiring,
            "amount": monthly_hiring_amount,
        },
        {
            "item_type": "taxable",
            "description": "Outstation Driver Batta",
            "quantity": outstation_duties,
            "rate": batta_rate,
            "amount": batta_amount,
        },
        {
            "item_type": "taxable",
            "description": "Extra Hours Charges",
            "quantity": extra_hours,
            "rate": extra_rate,
            "amount": extra_amount,
        },
        {
            "item_type": "taxable",
            "description": "Holiday Working Allowance",
            "quantity": holidays_worked,
            "rate": holiday_rate,
            "amount": holiday_amount,
        },
    ]

    # Fuel line item.
    # SAMRO uses Petrol; other profiles use Diesel.

    fuel_description = (
        "Petrol Charges" if is_samro else "Diesel Charges"
    )

    if fuel_amount:
        items.append({
            "item_type": "non-taxable",
            "description": fuel_description,
            "quantity": fuel_quantity,
            "rate": fuel_rate,
            "amount": fuel_amount,
        })

    # Engine Oil line item.
    if engine_oil:
        items.append({
            "item_type": "non-taxable",
            "description": "Engine Oil",
            "quantity": (
                total_km / 250 if not is_samro else 0
            ),
            "rate": 300,
            "amount": engine_oil,
        })

    # Parking and Toll line item.
    if parking_toll:
        items.append({
            "item_type": "non-taxable",
            "description": "Parking & Toll Charges",
            "quantity": 1,
            "rate": parking_toll,
            "amount": parking_toll,
        })

    # Driver Telephone Charges line item.
    if telephone_charges:
        items.append({
            "item_type": "non-taxable",
            "description": "Driver Telephone Charges",
            "quantity": 1,
            "rate": telephone_charges,
            "amount": telephone_charges,
        })

    # ---------------------------------------------------------
    # 8. TAXABLE AND NON-TAXABLE TOTALS
    # ---------------------------------------------------------

    # Taxable components:
    # Monthly Hiring + Outstation Batta
    # + Extra Hours + Holiday Allowance.

    taxable_subtotal = (
        monthly_hiring_amount
        + batta_amount
        + extra_amount
        + holiday_amount
    )

    # Non-taxable components:
    # Fuel + Engine Oil + Parking/Toll + Telephone.

    non_taxable_total = (
        fuel_amount
        + engine_oil
        + parking_toll
        + telephone_charges
    )

    # ---------------------------------------------------------
    # 9. GST CALCULATION
    # ---------------------------------------------------------

    cgst_rate = number("cgst_rate", 2.5)
    sgst_rate = number("sgst_rate", 2.5)

    # GST is calculated only on the taxable subtotal.

    cgst = taxable_subtotal * cgst_rate / 100
    sgst = taxable_subtotal * sgst_rate / 100

    # ---------------------------------------------------------
    # 10. GRAND TOTAL AND ROUND OFF
    # ---------------------------------------------------------

    # Grand Total is calculated before GST:
    # Taxable Subtotal + Non-Taxable Total.

    amount_before_rounding = (
        taxable_subtotal + non_taxable_total
    )

    # Round to the nearest whole rupee.

    grand_total = float(
        Decimal(str(amount_before_rounding)).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP
        )
    )

    round_off = grand_total - amount_before_rounding

    # ---------------------------------------------------------
    # 11. DATE VALIDATION
    # ---------------------------------------------------------

    start_date = optional_date("start_date")
    end_date = optional_date("end_date")

    if start_date and end_date and end_date < start_date:
        raise ValueError(
            "End Date cannot be earlier than Start Date"
        )

    # ---------------------------------------------------------
    # 12. RETURN DATA TO THE MONTHLY BILL API
    # ---------------------------------------------------------

    return {
        "invoice_number": str(
            data.get("invoice_number") or ""
        ).strip(),

        "invoice_date": parse_date(
            data.get("invoice_date"),
            "Invoice Date"
        ),

        "customer_name": str(
            data.get("customer_name") or ""
        ).strip(),

        "customer_profile": profile or None,

        "customer_address": str(
            data.get("customer_address") or ""
        ).strip(),

        "customer_gstin": str(
            data.get("customer_gstin") or ""
        ).strip(),

        "used_by": str(
            data.get("used_by") or ""
        ).strip() or None,

        "booked_by": str(
            data.get("booked_by") or ""
        ).strip(),

        "position": str(
            data.get("position") or ""
        ).strip() or None,

        "start_date": start_date,
        "end_date": end_date,

        "vehicle_number": str(
            data.get("vehicle_number") or ""
        ).strip(),

        "reference_number": str(
            data.get("reference_number") or ""
        ).strip(),

        # Monthly usage and fuel details.
        "total_km_driven": total_km,
        "km_per_liter": km_per_liter,
        "fuel_quantity": fuel_quantity,
        "fuel_rate": fuel_rate,

        "outstation_duties": outstation_duties,
        "extra_hours": extra_hours,
        "holidays_worked": holidays_worked,

        # Accounting totals.
        "taxable_subtotal": taxable_subtotal,

        "cgst_rate": cgst_rate,
        "cgst": cgst,

        "sgst_rate": sgst_rate,
        "sgst": sgst,

        "non_taxable_total": non_taxable_total,

        # Other charges.
        "engine_oil_amount": engine_oil,
        "parking_toll_amount": parking_toll,
        "telephone_charges_amount": telephone_charges,

        "round_off": round_off,
        "grand_total": grand_total,

        # Internal line items for PDF and accounting.
        "items": items,
    }




@app.post("/api/monthly-bills")
def create_monthly_bill():
    session = SessionLocal()

    try:
        data = request.get_json(silent=True) or {}

        values = calculate_monthly_bill_data(data)

        if not values["customer_name"]:
            raise ValueError("Customer Name is required")

        if not values["invoice_number"]:
            raise ValueError("Invoice Number is required")

        existing = session.execute(
            select(MonthlyBill.id).where(
                MonthlyBill.invoice_number == values["invoice_number"]
            )
        ).scalar_one_or_none()

        if existing is not None:
            return error_response(
                "Monthly Bill invoice number already exists",
                409,
                "DUPLICATE_MONTHLY_BILL_NUMBER",
            )

        items = values.pop("items")

        bill = MonthlyBill(**values)

        for item_data in items:
            bill.items.append(MonthlyBillItem(**item_data))

        session.add(bill)
        session.commit()

        refreshed = session.execute(
            select(MonthlyBill)
            .options(joinedload(MonthlyBill.items))
            .where(MonthlyBill.id == bill.id)
        ).unique().scalar_one()

        return jsonify({
            "success": True,
            "message": "Monthly Bill created successfully",
            "monthly_bill": serialize_monthly_bill(refreshed),
        }), 201

    except ValueError as exc:
        session.rollback()
        return error_response(str(exc), 400, "VALIDATION_ERROR")

    except IntegrityError:
        session.rollback()
        return error_response(
            "Monthly Bill invoice number already exists",
            409,
            "DUPLICATE_MONTHLY_BILL_NUMBER",
        )

    except Exception as exc:
        session.rollback()
        traceback.print_exc()
        return error_response(
            "Unable to save Monthly Bill",
            500,
            "SERVER_ERROR",
            str(exc),
        )

    finally:
        session.close()

@app.get("/api/monthly-bills/<int:bill_id>/pdf")
def monthly_bill_pdf(bill_id):
    session = SessionLocal()

    try:
        bill = (
            session.execute(
                select(MonthlyBill)
                .where(MonthlyBill.id == bill_id)
            )
            .scalar_one_or_none()
        )

        if not bill:
            return error_response(
                "Monthly Bill not found",
                404,
                "NOT_FOUND"
            )
        from services import build_monthly_bill_pdf

        pdf = build_monthly_bill_pdf(bill)

        filename = (
            f"{bill.invoice_number.replace('/', '-')}.pdf"
        )

        return send_file(
            __import__("io").BytesIO(pdf),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename
        )

    except Exception as exc:
        traceback.print_exc()

        return error_response(
            "Unable to generate Monthly Bill PDF",
            500,
            "MONTHLY_BILL_PDF_ERROR",
            str(exc)
        )

    finally:
        session.close()

@app.get("/api/monthly-bills")

def list_monthly_bills():
    session = SessionLocal()

    try:
        month_value = (request.args.get("month") or "").strip()
        year_value = (request.args.get("year") or "").strip()

        query = (
            select(MonthlyBill)
            .options(joinedload(MonthlyBill.items))
            .order_by(
                MonthlyBill.invoice_date.desc(),
                MonthlyBill.id.desc()
            )
        )

        if month_value and year_value:
            try:
                selected_month = int(month_value)
                selected_year = int(year_value)

                if selected_month < 1 or selected_month > 12:
                    return error_response(
                        "Month must be between 1 and 12",
                        400,
                        "INVALID_MONTH"
                    )

                if selected_year < 2000 or selected_year > 2100:
                    return error_response(
                        "Invalid year",
                        400,
                        "INVALID_YEAR"
                    )

                start_date = date(selected_year, selected_month, 1)

                if selected_month == 12:
                    next_month_date = date(selected_year + 1, 1, 1)
                else:
                    next_month_date = date(
                        selected_year,
                        selected_month + 1,
                        1
                    )

                query = query.where(
                    MonthlyBill.invoice_date >= start_date,
                    MonthlyBill.invoice_date < next_month_date
                )

            except ValueError:
                return error_response(
                    "Month and year must be valid numbers",
                    400,
                    "INVALID_MONTH_YEAR"
                )

        bills = (
            session.execute(query)
            .unique()
            .scalars()
            .all()
        )

        return jsonify({
            "success": True,
            "monthly_bills": [
                serialize_monthly_bill(bill)
                for bill in bills
            ]
        })

    except Exception as exc:
        session.rollback()
        traceback.print_exc()

        return error_response(
            "Unable to load Monthly Bills",
            500,
            "MONTHLY_BILL_LIST_ERROR",
            str(exc)
        )

    finally:
        session.close()

@app.get("/api/monthly-bills/<int:bill_id>")
def get_monthly_bill(bill_id):
    session = SessionLocal()

    try:
        bill = (
            session.execute(
                select(MonthlyBill)
                .options(
                    joinedload(MonthlyBill.items)
                )
                .where(
                    MonthlyBill.id == bill_id
                )
            )
            .unique()
            .scalar_one_or_none()
        )

        if not bill:
            return error_response(
                "Monthly Bill not found",
                404,
                "NOT_FOUND"
            )

        return jsonify({
            "success": True,
            "monthly_bill": serialize_monthly_bill(bill)
        })

    except Exception as exc:
        session.rollback()
        traceback.print_exc()

        return error_response(
            "Unable to load Monthly Bill",
            500,
            "MONTHLY_BILL_GET_ERROR",
            str(exc)
        )

    finally:
        session.close()

@app.put("/api/monthly-bills/<int:bill_id>")
def update_monthly_bill(bill_id):
    session = SessionLocal()

    try:
        bill = session.execute(
            select(MonthlyBill)
            .options(joinedload(MonthlyBill.items))
            .where(MonthlyBill.id == bill_id)
        ).unique().scalar_one_or_none()

        if not bill:
            return error_response(
                "Monthly Bill not found", 404, "NOT_FOUND"
            )

        data = request.get_json(silent=True) or {}
        values = calculate_monthly_bill_data(data)

        if not values["customer_name"]:
            raise ValueError("Customer Name is required")

        if not values["invoice_number"]:
            raise ValueError("Invoice Number is required")

        existing = session.execute(
            select(MonthlyBill.id).where(
                MonthlyBill.invoice_number == values["invoice_number"],
                MonthlyBill.id != bill_id,
            )
        ).scalar_one_or_none()

        if existing is not None:
            return error_response(
                "Monthly Bill invoice number already exists",
                409,
                "DUPLICATE_MONTHLY_BILL_NUMBER",
            )

        items = values.pop("items")

        for key, value in values.items():
            setattr(bill, key, value)

        bill.items.clear()
        session.flush()

        for item_data in items:
            bill.items.append(MonthlyBillItem(**item_data))

        session.commit()

        refreshed = session.execute(
            select(MonthlyBill)
            .options(joinedload(MonthlyBill.items))
            .where(MonthlyBill.id == bill_id)
        ).unique().scalar_one()

        return jsonify({
            "success": True,
            "message": "Monthly Bill updated successfully",
            "monthly_bill": serialize_monthly_bill(refreshed),
        })

    except ValueError as exc:
        session.rollback()
        return error_response(str(exc), 400, "VALIDATION_ERROR")

    except IntegrityError:
        session.rollback()
        return error_response(
            "Monthly Bill invoice number already exists",
            409,
            "DUPLICATE_MONTHLY_BILL_NUMBER",
        )

    except Exception as exc:
        session.rollback()
        traceback.print_exc()
        return error_response(
            "Unable to update Monthly Bill",
            500,
            "MONTHLY_BILL_UPDATE_ERROR",
            str(exc),
        )

    finally:
        session.close()

